# -*- coding: utf-8 -*-
import json
import os
import re
import shutil
import threading
import uuid
import urllib.error
import urllib.request
import warnings
from copy import deepcopy
from datetime import datetime
from typing import Any, Dict, Iterable, List, Optional, Tuple

from django.conf import settings
from django.db import close_old_connections, transaction
from django.utils import timezone
from openpyxl import load_workbook

from task.models import ListingOptimizationJob, ListingOptimizationRow


TITLE_KEY = "Product Name（标题）"
DESCRIPTION_KEY = "Product Description（描述）"
BULLET_1_KEY = "Key Product Features 1（五点1）"
BULLET_2_KEY = "Key Product Features 2（五点2）"
BULLET_3_KEY = "Key Product Features 3（五点3）"
BULLET_4_KEY = "Key Product Features 4（五点4）"
BULLET_5_KEY = "Key Product Features 5（五点5）"
SEARCH_TERMS_KEY = "Search Terms（后台关键词）"
MAIN_IMAGE_KEY = "Main Image URL（主图）"

TRANSLATABLE_KEYS = [
    TITLE_KEY,
    DESCRIPTION_KEY,
    BULLET_1_KEY,
    BULLET_2_KEY,
    BULLET_3_KEY,
    BULLET_4_KEY,
    BULLET_5_KEY,
    SEARCH_TERMS_KEY,
]
REQUIRED_LISTING_KEYS = [*TRANSLATABLE_KEYS, MAIN_IMAGE_KEY]

OUTPUT_KEYS: List[str] = []
for _field_key in TRANSLATABLE_KEYS:
    OUTPUT_KEYS.extend([_field_key, f"{_field_key}中文"])
OUTPUT_KEYS.append(MAIN_IMAGE_KEY)
del _field_key

EXCEL_COLUMN_MAP = {
    "item_name": TITLE_KEY,
    "product_description": DESCRIPTION_KEY,
    "bullet_point1": BULLET_1_KEY,
    "bullet_point2": BULLET_2_KEY,
    "bullet_point3": BULLET_3_KEY,
    "bullet_point4": BULLET_4_KEY,
    "bullet_point5": BULLET_5_KEY,
    "generic_keywords": SEARCH_TERMS_KEY,
    "main_image_url": MAIN_IMAGE_KEY,
}

HEADER_ROW = 3
DATA_START_ROW = 4
MAX_UPLOAD_SIZE = 10 * 1024 * 1024

_active_jobs = set()
_worker_lock = threading.Lock()


class ListingValidationError(ValueError):
    pass


class DashScopeError(RuntimeError):
    def __init__(self, message: str, status: Optional[int] = None, body: str = "") -> None:
        super().__init__(message)
        self.status = status
        self.body = body


class QwenAmazonListingOptimizer:
    DEFAULT_PROMPT = """
你是拥有20年美国 Amazon POD 产品运营经验的资深卖家，擅长为美亚买家编写自然、合规、可转化的 listing。
你熟悉美国搜索习惯，并以 A10 相关性、COSMO/Rufus 语义理解、GEO 生成式检索友好表达为目标：用清晰的“产品词 + 主题词 + 属性 + 场景 + 人群”线性结构组织文案。

任务：
基于输入的主题、产品属性、图片 URL 和现有文案，重写 Amazon listing。输出必须包含英文标题、五点描述、产品描述、后台搜索关键词、这些英文内容的中文意思，以及主图 URL。

硬性规则：
1. 只输出一个 JSON 对象，不要 Markdown，不要解释，不要多余字段。
2. JSON 的 key 必须严格等于用户提供的输出字段清单；顺序也尽量保持一致。
3. 除主图 URL 外，每个英文 listing 字段后面都必须增加对应的“中文”字段，例如 "Key Product Features 1（五点1）中文"。
4. “中文”字段必须是对应英文内容的自然中文翻译或中文意思，不要写优化说明，不要逐条解释你为什么这样写。
5. 标题使用地道美式英语，不超过 125 个字符；前半段放核心产品词与主题词，避免堆砌和重复。
6. 五点描述共 5 条，每条只写一段英文；围绕主题自然埋词，并分别覆盖材质/舒适度、版型/尺寸、图案风格、穿搭场景、送礼/日常使用。
7. 产品描述写成流畅英文自然段，突出 POD 图案、复古水洗棉、低帽冠、可调节范围、男女通用和使用场景。
8. 后台搜索关键词不超过 250 个字符；全部小写；用空格分隔；不重复；不要标点、品牌名、ASIN、竞品词或侵权词。
9. “Search Terms（后台关键词）中文”按正常中文意思翻译即可，可以是自然短语，不需要遵守后台关键词的空格格式。
10. 侵权与合规过滤：不要使用商标、品牌、球队、电影、明星、角色、歌词、组织名称等受保护内容；不要写 official、licensed、authentic、best、#1、guaranteed、medical claims 等高风险表达。
11. 不要虚构认证、销量、库存、物流、环保认证或“latest 2026 release”等无法从输入确认的事实。
12. 产品是复古水洗棉、低帽冠、可调节、男女通用的 dad hat/baseball cap；不要把产品误写成 shirt。
13. 主图 URL 必须原样返回，并且不要为主图 URL 添加中文字段。
14. 输入字典中如果有额外字段，可以作为上下文参考，但不要把额外字段返回。
"""

    def __init__(
        self,
        api_key: Optional[str] = None,
        model: Optional[str] = None,
        endpoint: Optional[str] = None,
        timeout: int = 90,
    ) -> None:
        self.api_key = api_key or getattr(settings, "DASHSCOPE_API_KEY", "") or os.getenv("DASHSCOPE_API_KEY", "")
        self.model = model or getattr(settings, "DASHSCOPE_MODEL", "") or "qwen-turbo-latest"
        self.endpoint = (
            endpoint
            or getattr(settings, "DASHSCOPE_ENDPOINT", "")
            or "https://dashscope.aliyuncs.com/compatible-mode/v1/chat/completions"
        )
        self.timeout = timeout
        if not self.api_key:
            raise DashScopeError("缺少 DashScope API key，请在 .env 中配置 DASHSCOPE_API_KEY。")

    def optimize(self, listing: Dict[str, Any]) -> Dict[str, str]:
        source_listing = validate_listing_dict(listing)
        last_error: Optional[DashScopeError] = None
        for candidate_model in self._candidate_models():
            for include_response_format in (True, False):
                for include_thinking_flag in (True, False):
                    payload = self._build_payload(
                        source_listing,
                        candidate_model,
                        include_response_format,
                        include_thinking_flag,
                    )
                    try:
                        content = self._chat_completion(payload)
                        parsed = self._extract_json_object(content)
                        return normalize_optimizer_result(parsed, source_listing)
                    except DashScopeError as exc:
                        last_error = exc
                        if self._should_retry(exc):
                            continue
                        raise

        if last_error:
            raise last_error
        raise DashScopeError("千问接口调用失败，但未返回具体错误。")

    def _candidate_models(self) -> Iterable[str]:
        fallbacks = (self.model, "qwen-turbo-latest", "qwen-flash", "qwen-plus")
        seen = set()
        for model_name in fallbacks:
            if model_name and model_name not in seen:
                seen.add(model_name)
                yield model_name

    def _build_payload(
        self,
        listing: Dict[str, str],
        model: str,
        include_response_format: bool,
        include_thinking_flag: bool,
    ) -> Dict[str, Any]:
        user_prompt = (
            "请先在内部分析主题、产品属性、美国买家搜索词和潜在侵权风险，"
            "然后只返回最终 JSON 对象。\n\n"
            f"必须返回的字段清单：\n{json.dumps(OUTPUT_KEYS, ensure_ascii=False, indent=2)}\n\n"
            f"输入 listing：\n{json.dumps(listing, ensure_ascii=False, indent=2)}"
        )
        payload: Dict[str, Any] = {
            "model": model,
            "messages": [
                {"role": "system", "content": self.DEFAULT_PROMPT},
                {"role": "user", "content": user_prompt},
            ],
            "temperature": 0.35,
            "max_tokens": 2200,
        }
        if include_response_format:
            payload["response_format"] = {"type": "json_object"}
        if include_thinking_flag:
            payload["enable_thinking"] = False
        return payload

    def _chat_completion(self, payload: Dict[str, Any]) -> str:
        data = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        request = urllib.request.Request(
            self.endpoint,
            data=data,
            headers={
                "Authorization": f"Bearer {self.api_key}",
                "Content-Type": "application/json",
            },
            method="POST",
        )

        try:
            with urllib.request.urlopen(request, timeout=self.timeout) as response:
                raw = response.read().decode("utf-8")
        except urllib.error.HTTPError as exc:
            body = exc.read().decode("utf-8", errors="replace")
            raise DashScopeError(f"DashScope API HTTP {exc.code}: {body}", exc.code, body) from exc
        except urllib.error.URLError as exc:
            raise DashScopeError(f"DashScope API request failed: {exc.reason}") from exc

        try:
            result = json.loads(raw)
            return result["choices"][0]["message"]["content"]
        except (KeyError, IndexError, TypeError, json.JSONDecodeError) as exc:
            raise DashScopeError(f"无法解析 DashScope 返回内容: {raw[:500]}") from exc

    @staticmethod
    def _extract_json_object(content: str) -> Dict[str, Any]:
        text = content.strip()
        if text.startswith("```"):
            text = re.sub(r"^```(?:json)?\s*", "", text, flags=re.IGNORECASE)
            text = re.sub(r"\s*```$", "", text)

        try:
            parsed = json.loads(text)
        except json.JSONDecodeError:
            match = re.search(r"\{.*\}", text, flags=re.DOTALL)
            if not match:
                raise DashScopeError(f"模型没有返回 JSON 对象: {content[:500]}")
            parsed = json.loads(match.group(0))

        if not isinstance(parsed, dict):
            raise DashScopeError("模型返回的 JSON 不是对象。")
        return parsed

    @staticmethod
    def _should_retry(error: DashScopeError) -> bool:
        if error.status != 400:
            return False
        retry_markers = (
            "enable_thinking",
            "model",
            "parameter",
            "invalid",
            "not support",
            "not exist",
            "unsupported",
            "response_format",
        )
        text = f"{error} {error.body}".lower()
        return any(marker in text for marker in retry_markers)


def validate_listing_dict(listing: Dict[str, Any]) -> Dict[str, str]:
    if not isinstance(listing, dict):
        raise ListingValidationError("listing 必须是 dict。")

    missing_keys = [key for key in REQUIRED_LISTING_KEYS if key not in listing]
    if missing_keys:
        raise ListingValidationError(f"listing 缺少必填字段: {', '.join(missing_keys)}")

    normalized = deepcopy(listing)
    for key in REQUIRED_LISTING_KEYS:
        value = normalized.get(key)
        if value is None:
            normalized[key] = ""
        elif not isinstance(value, str):
            normalized[key] = str(value)
    return normalized


def normalize_optimizer_result(result: Dict[str, Any], original: Dict[str, str]) -> Dict[str, str]:
    missing_keys = [key for key in OUTPUT_KEYS if key not in result]
    if missing_keys:
        raise DashScopeError(f"模型返回结果缺少字段: {', '.join(missing_keys)}")

    normalized: Dict[str, str] = {}
    for key in OUTPUT_KEYS:
        if key == MAIN_IMAGE_KEY:
            normalized[key] = clean_text(original[MAIN_IMAGE_KEY])
        else:
            normalized[key] = clean_text(result[key])

    normalized[TITLE_KEY] = clip_at_word(normalized[TITLE_KEY], 125)
    normalized[SEARCH_TERMS_KEY] = clip_at_word(clean_search_terms(normalized[SEARCH_TERMS_KEY]), 250)
    return normalized


def clean_text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def clean_search_terms(value: str) -> str:
    tokens = re.findall(r"[a-z0-9][a-z0-9'-]*", value.lower())
    unique_tokens = []
    seen = set()
    for token in tokens:
        token = token.strip("'-")
        if token and token not in seen:
            seen.add(token)
            unique_tokens.append(token)
    return " ".join(unique_tokens)


def clip_at_word(value: str, limit: int) -> str:
    value = clean_text(value)
    if len(value) <= limit:
        return value
    clipped = value[:limit].rstrip()
    if " " in clipped:
        clipped = clipped.rsplit(" ", 1)[0]
    return clipped.rstrip(" ,.;:-")


def save_uploaded_listing_file(uploaded_file) -> Tuple[str, str]:
    if not uploaded_file.name.lower().endswith((".xlsx", ".xlsm")):
        raise ListingValidationError("仅支持 Excel 文件（.xlsx、.xlsm）。")
    if uploaded_file.size > MAX_UPLOAD_SIZE:
        raise ListingValidationError("文件大小不能超过 10MB。")

    temp_dir = os.path.join(settings.MEDIA_ROOT, "temp_uploads", "listing_optimizer", str(uuid.uuid4()))
    os.makedirs(temp_dir, exist_ok=True)
    file_path = os.path.join(temp_dir, uploaded_file.name)
    with open(file_path, "wb+") as destination:
        for chunk in uploaded_file.chunks():
            destination.write(chunk)

    relative_path = os.path.relpath(file_path, settings.MEDIA_ROOT)
    return relative_path.replace("\\", "/"), file_path


def create_job_from_excel(uploaded_file, user) -> ListingOptimizationJob:
    relative_path, full_path = save_uploaded_listing_file(uploaded_file)
    rows = read_listing_rows(full_path)

    with transaction.atomic():
        job = ListingOptimizationJob.objects.create(
            created_by=user,
            original_filename=uploaded_file.name,
            original_file_path=relative_path,
            status=ListingOptimizationJob.STATUS_PENDING,
            total_rows=len(rows),
        )
        ListingOptimizationRow.objects.bulk_create(
            [
                ListingOptimizationRow(
                    job=job,
                    row_number=row_number,
                    source_data=source_data,
                    optimized_data={},
                    status=ListingOptimizationRow.STATUS_PENDING,
                )
                for row_number, source_data in rows
            ]
        )
    return job


def read_listing_rows(full_path: str) -> List[Tuple[int, Dict[str, str]]]:
    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="Workbook contains no default style, apply openpyxl's default",
            category=UserWarning,
            module="openpyxl.styles.stylesheet",
        )
        workbook = load_workbook(full_path, read_only=True, data_only=True)
    worksheet = workbook.active
    header_map = _get_header_map(worksheet)
    missing_columns = [column for column in EXCEL_COLUMN_MAP if column not in header_map]
    if missing_columns:
        raise ListingValidationError(f"Excel 缺少必需列: {', '.join(missing_columns)}")

    rows = []
    for excel_row in range(DATA_START_ROW, worksheet.max_row + 1):
        title_value = worksheet.cell(excel_row, header_map["item_name"]).value
        if clean_text(title_value) == "":
            continue

        source_data = {}
        for excel_column, listing_key in EXCEL_COLUMN_MAP.items():
            value = worksheet.cell(excel_row, header_map[excel_column]).value
            source_data[listing_key] = clean_text(value)
        rows.append((excel_row, source_data))

    workbook.close()
    if not rows:
        raise ListingValidationError("Excel 中没有可优化的 listing 行。")
    return rows


def _get_header_map(worksheet) -> Dict[str, int]:
    header_map = {}
    for cell in worksheet[HEADER_ROW]:
        value = clean_text(cell.value)
        if value:
            header_map[value] = cell.column
    return header_map


def start_listing_optimization_job(job_id: int) -> bool:
    with _worker_lock:
        if job_id in _active_jobs:
            return False
        _active_jobs.add(job_id)

    worker = threading.Thread(
        target=_run_listing_optimization_job,
        args=(job_id,),
        daemon=True,
        name=f"listing-optimizer-{job_id}",
    )
    worker.start()
    return True


def ensure_listing_job_running(job: ListingOptimizationJob) -> None:
    if job.status in {ListingOptimizationJob.STATUS_PENDING, ListingOptimizationJob.STATUS_PROCESSING}:
        start_listing_optimization_job(job.id)


def pause_listing_optimization_job(job: ListingOptimizationJob) -> ListingOptimizationJob:
    if job.status in {ListingOptimizationJob.STATUS_PENDING, ListingOptimizationJob.STATUS_PROCESSING}:
        job.status = ListingOptimizationJob.STATUS_PAUSED
        job.save(update_fields=["status", "updated_at"])
    return job


def _run_listing_optimization_job(job_id: int) -> None:
    close_old_connections()
    try:
        job = ListingOptimizationJob.objects.get(id=job_id)
        ListingOptimizationRow.objects.filter(
            job=job,
            status=ListingOptimizationRow.STATUS_PROCESSING,
        ).update(status=ListingOptimizationRow.STATUS_PENDING, error_message="")

        if job.status == ListingOptimizationJob.STATUS_PAUSED:
            return

        if job.status != ListingOptimizationJob.STATUS_PROCESSING:
            job.status = ListingOptimizationJob.STATUS_PROCESSING
            if not job.started_at:
                job.started_at = timezone.now()
            job.save(update_fields=["status", "started_at", "updated_at"])

        optimizer = QwenAmazonListingOptimizer()
        pending_rows = ListingOptimizationRow.objects.filter(
            job_id=job_id,
            status=ListingOptimizationRow.STATUS_PENDING,
        ).order_by("row_number")

        for row in pending_rows:
            current_status = ListingOptimizationJob.objects.filter(id=job_id).values_list("status", flat=True).first()
            if current_status == ListingOptimizationJob.STATUS_PAUSED:
                return

            row.status = ListingOptimizationRow.STATUS_PROCESSING
            row.error_message = ""
            row.save(update_fields=["status", "error_message", "updated_at"])
            try:
                optimized = optimizer.optimize(row.source_data)
                row.optimized_data = optimized
                row.status = ListingOptimizationRow.STATUS_COMPLETED
                row.error_message = ""
                row.save(update_fields=["optimized_data", "status", "error_message", "updated_at"])
            except Exception as exc:
                row.status = ListingOptimizationRow.STATUS_FAILED
                row.error_message = str(exc)[:2000]
                row.save(update_fields=["status", "error_message", "updated_at"])
            _refresh_job_progress(job_id)

            current_status = ListingOptimizationJob.objects.filter(id=job_id).values_list("status", flat=True).first()
            if current_status == ListingOptimizationJob.STATUS_PAUSED:
                return

        _finish_job(job_id)
    except Exception as exc:
        ListingOptimizationJob.objects.filter(id=job_id).update(
            status=ListingOptimizationJob.STATUS_FAILED,
            error_message=str(exc)[:2000],
            completed_at=timezone.now(),
        )
    finally:
        with _worker_lock:
            _active_jobs.discard(job_id)
        close_old_connections()


def _refresh_job_progress(job_id: int) -> None:
    optimized_count = ListingOptimizationRow.objects.filter(
        job_id=job_id,
        status=ListingOptimizationRow.STATUS_COMPLETED,
    ).count()
    failed_count = ListingOptimizationRow.objects.filter(
        job_id=job_id,
        status=ListingOptimizationRow.STATUS_FAILED,
    ).count()
    ListingOptimizationJob.objects.filter(id=job_id).update(
        optimized_rows=optimized_count,
        failed_rows=failed_count,
        updated_at=timezone.now(),
    )


def _finish_job(job_id: int) -> None:
    _refresh_job_progress(job_id)
    job = ListingOptimizationJob.objects.get(id=job_id)
    has_failed = job.failed_rows > 0
    ListingOptimizationJob.objects.filter(id=job_id).update(
        status=ListingOptimizationJob.STATUS_FAILED if has_failed else ListingOptimizationJob.STATUS_COMPLETED,
        completed_at=timezone.now(),
        updated_at=timezone.now(),
    )


def update_row_optimized_data(row: ListingOptimizationRow, optimized_data: Dict[str, Any]) -> ListingOptimizationRow:
    cleaned = {}
    base_data = row.optimized_data or row.source_data or {}
    for key in OUTPUT_KEYS:
        if key == MAIN_IMAGE_KEY:
            cleaned[key] = row.source_data.get(MAIN_IMAGE_KEY, "")
        else:
            cleaned[key] = clean_text(optimized_data.get(key, base_data.get(key, "")))
    cleaned[TITLE_KEY] = clip_at_word(cleaned[TITLE_KEY], 125)
    cleaned[SEARCH_TERMS_KEY] = clip_at_word(clean_search_terms(cleaned[SEARCH_TERMS_KEY]), 250)
    row.optimized_data = cleaned
    row.save(update_fields=["optimized_data", "updated_at"])
    return row


def generate_optimized_excel(job: ListingOptimizationJob) -> str:
    original_full_path = os.path.join(settings.MEDIA_ROOT, job.original_file_path)
    if not os.path.exists(original_full_path):
        raise FileNotFoundError("原始上传文件不存在，无法生成下载文件。")

    output_dir = os.path.join(settings.MEDIA_ROOT, "temp_uploads", "listing_optimizer_outputs", str(job.id))
    os.makedirs(output_dir, exist_ok=True)
    filename_root, filename_ext = os.path.splitext(job.original_filename)
    output_name = f"{filename_root}_listing优化_{datetime.now().strftime('%Y%m%d%H%M%S')}{filename_ext or '.xlsx'}"
    output_full_path = os.path.join(output_dir, output_name)
    shutil.copy2(original_full_path, output_full_path)

    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="Workbook contains no default style, apply openpyxl's default",
            category=UserWarning,
            module="openpyxl.styles.stylesheet",
        )
        workbook = load_workbook(output_full_path)
    worksheet = workbook.active
    header_map = _get_header_map(worksheet)
    missing_columns = [column for column in EXCEL_COLUMN_MAP if column not in header_map]
    if missing_columns:
        workbook.close()
        raise ListingValidationError(f"Excel 缺少必需列: {', '.join(missing_columns)}")

    rows = job.rows.all().order_by("row_number")
    reverse_map = {listing_key: excel_column for excel_column, listing_key in EXCEL_COLUMN_MAP.items()}
    for row in rows:
        data = row.optimized_data or row.source_data
        for listing_key in REQUIRED_LISTING_KEYS:
            excel_column_name = reverse_map[listing_key]
            worksheet.cell(row.row_number, header_map[excel_column_name]).value = data.get(listing_key, "")

    workbook.save(output_full_path)
    workbook.close()

    relative_output_path = os.path.relpath(output_full_path, settings.MEDIA_ROOT).replace("\\", "/")
    job.output_file_path = relative_output_path
    job.save(update_fields=["output_file_path", "updated_at"])
    return relative_output_path


def serialize_job(job: ListingOptimizationJob, include_rows: bool = False) -> Dict[str, Any]:
    data = {
        "id": job.id,
        "original_filename": job.original_filename,
        "status": job.status,
        "status_display": job.get_status_display(),
        "total_rows": job.total_rows,
        "optimized_rows": job.optimized_rows,
        "failed_rows": job.failed_rows,
        "error_message": job.error_message,
        "output_file_path": job.output_file_path,
        "created_at": job.created_at.strftime("%Y-%m-%d %H:%M:%S") if job.created_at else "",
        "updated_at": job.updated_at.strftime("%Y-%m-%d %H:%M:%S") if job.updated_at else "",
        "download_url": "",
    }
    if job.output_file_path:
        data["download_url"] = f"/api/tasks/listing-optimizer/jobs/{job.id}/download/"
    if include_rows:
        data["rows"] = [serialize_row(row) for row in job.rows.all().order_by("row_number")]
    return data


def serialize_row(row: ListingOptimizationRow) -> Dict[str, Any]:
    display_number = max((row.row_number or DATA_START_ROW) - DATA_START_ROW + 1, 1)
    return {
        "id": row.id,
        "row_number": row.row_number,
        "display_number": display_number,
        "status": row.status,
        "status_display": row.get_status_display(),
        "source_data": row.source_data or {},
        "optimized_data": row.optimized_data or {},
        "error_message": row.error_message,
        "updated_at": row.updated_at.strftime("%Y-%m-%d %H:%M:%S") if row.updated_at else "",
    }
