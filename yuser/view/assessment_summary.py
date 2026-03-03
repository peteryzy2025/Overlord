# yuser/view/assessment_summary.py

import io
from decimal import Decimal

from django.contrib.auth.decorators import login_required
from django.http import JsonResponse, HttpResponse
from django.shortcuts import render
from django.utils.decorators import method_decorator
from django.views import View
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from yuser.view_assessment_management import AssessmentManagementView
from yuser.models import PerformanceAssessment


@method_decorator(login_required, name='dispatch')
class AssessmentSummaryView(AssessmentManagementView, View):
    """考核成绩汇总页面"""

    def get_weight(self, assess_type):
        """获取权重配置"""
        if assess_type == 'temu_operation':
            return Decimal('0.7'), Decimal('0.3')
        elif assess_type == 'amazon_operation':
            return Decimal('0.6'), Decimal('0.4')
        else:  # amazon_assistant
            return Decimal('0.3'), Decimal('0.7')

    def calculate_weighted_scores(self, assessment):
        """按详情页同口径拆分三项分数"""
        kpi_weight, behavior_weight = self.get_weight(assessment.assess_type)
        details_dict = {d.item_key: Decimal(str(d.score_value)) for d in assessment.score_details.all()}

        # 与 /assessment/<int>/ 一致：业绩分由 achievement_rate 动态计算，而不是读取 detail.performance_score
        performance_raw = Decimal(str(self.calculate_performance_score(assessment)))
        accident_score = details_dict.get('accident_score', Decimal('0'))

        behavior_raw = sum(
            score for key, score in details_dict.items()
            if key not in {'performance_score', 'accident_score', 'extra_bonus', 'shop_activation'}
        )

        extra_bonus = details_dict.get('extra_bonus', Decimal('0'))
        shop_activation = details_dict.get('shop_activation', Decimal('0'))
        bonus_score = extra_bonus + shop_activation

        performance_weighted = (performance_raw + accident_score) * kpi_weight
        behavior_weighted = behavior_raw * behavior_weight

        return {
            'performance_score': round(float(performance_weighted), 2),
            'behavior_score': round(float(behavior_weighted), 2),
            'bonus_score': round(float(bonus_score), 2),
        }

    def _get_assessments_queryset(self, request, month=''):
        assessments = PerformanceAssessment.objects.select_related(
            'employee', 'employee__operational_account'
        ).prefetch_related('score_details').filter(
            employee__company=request.user.company,
            status=PerformanceAssessment.STATUS_CONFIRMED
        )

        if month:
            assessments = assessments.filter(month=month)

        return assessments.order_by(
            'employee__operational_account__ops_group',
            'employee__first_name',
            'assess_type'
        )

    def _group_assessments(self, assessments):
        grouped_data = {}
        for assessment in assessments:
            ops_group = assessment.employee.operational_account.ops_group if hasattr(
                assessment.employee, 'operational_account') and assessment.employee.operational_account else '未分组'

            if ops_group not in grouped_data:
                grouped_data[ops_group] = {'leader': None, 'members': []}

            if assessment.employee.role == '运营组长':
                grouped_data[ops_group]['leader'] = assessment
            else:
                grouped_data[ops_group]['members'].append(assessment)
        return grouped_data

    def _build_summary_rows(self, grouped_data):
        data_list = []
        for ops_group, group_data in grouped_data.items():
            leader = group_data['leader']
            members = group_data['members']

            group_target = sum(m.target_orders for m in members) if members else 0
            group_actual = sum(m.actual_orders for m in members) if members else 0
            group_rate = (group_actual / group_target * 100) if group_target > 0 else 0
            group_achieved = '是' if group_rate >= 100 else '否'

            data_list.append({
                'id': None,
                'is_leader_row': True,
                'employee_name': leader.employee.first_name if leader else ops_group,
                'role': '运营组长',
                'ops_group': ops_group,
                'assess_type_display': '-',
                'final_score': '-',
                'target_orders': group_target,
                'actual_orders': group_actual,
                'achievement_rate': round(group_rate, 2),
                'is_achieved': group_achieved,
                'performance_score': '-',
                'behavior_score': '-',
                'bonus_score': '-',
            })

            for assessment in members:
                scores = self.calculate_weighted_scores(assessment)
                is_achieved = '是' if assessment.actual_orders >= assessment.target_orders else '否'

                data_list.append({
                    'id': assessment.id,
                    'is_leader_row': False,
                    'employee_name': assessment.employee.first_name,
                    'role': assessment.employee.role or '-',
                    'ops_group': ops_group,
                    'assess_type_display': assessment.get_assess_type_display(),
                    'final_score': assessment.final_score or 0,
                    'target_orders': assessment.target_orders,
                    'actual_orders': assessment.actual_orders,
                    'achievement_rate': assessment.achievement_rate or 0,
                    'is_achieved': is_achieved,
                    'performance_score': scores['performance_score'],
                    'behavior_score': scores['behavior_score'],
                    'bonus_score': scores['bonus_score'],
                })

            if leader and leader.employee.role == '运营组长':
                leader_in_members = any(m.id == leader.id for m in members)
                if not leader_in_members:
                    scores = self.calculate_weighted_scores(leader)
                    is_achieved = '是' if leader.actual_orders >= leader.target_orders else '否'
                    data_list.append({
                        'id': leader.id,
                        'is_leader_row': False,
                        'employee_name': leader.employee.first_name + '（个人）',
                        'role': leader.employee.role or '-',
                        'ops_group': ops_group,
                        'assess_type_display': leader.get_assess_type_display(),
                        'final_score': leader.final_score or 0,
                        'target_orders': leader.target_orders,
                        'actual_orders': leader.actual_orders,
                        'achievement_rate': leader.achievement_rate or 0,
                        'is_achieved': is_achieved,
                        'performance_score': scores['performance_score'],
                        'behavior_score': scores['behavior_score'],
                        'bonus_score': scores['bonus_score'],
                    })
        return data_list

    def get(self, request):
        """渲染汇总页面"""
        # 权限检查：只有555管理员可以访问
        if not request.user.permission_configs.filter(code=555).exists():
            return JsonResponse({'success': False, 'message': '无权限访问'}, status=403)

        month = request.GET.get('month', '')

        assessments = self._get_assessments_queryset(request, month=month)
        grouped_data = self._group_assessments(assessments)
        data_list = self._build_summary_rows(grouped_data)

        context = {
            'active_page': 'assessment_summary',
            'active_nav': 'management',
            'month': month,
            'data_list': data_list,
            'is_admin': True,
        }

        return render(request, 'assessment_summary.html', context)

    def post(self, request):
        """处理Excel导出"""
        from decimal import Decimal  # 放这里或文件顶部都行

        # 权限检查：只有555管理员可以导出
        if not request.user.permission_configs.filter(code=555).exists():
            return JsonResponse({'success': False, 'message': '无权限导出'}, status=403)

        month = request.POST.get('month', '')
        if not month:
            return JsonResponse({'success': False, 'message': '请选择考核月份'}, status=400)

        # 解析年月用于文件名和标题
        year, mon = month.split('-')
        file_name = f"{year}年{mon}月运营部kpi考核成绩汇总.xlsx"
        big_title = f"{year}年{mon}月考核成绩汇总"

        assessments = self._get_assessments_queryset(request, month=month)
        grouped_data = self._group_assessments(assessments)
        data_list = self._build_summary_rows(grouped_data)

        wb = Workbook()
        ws = wb.active
        ws.title = f"{year}年{mon}月KPI汇总"

        # 定义样式
        title_font = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF')
        header_font = Font(name='微软雅黑', size=10, bold=True, color='2d3748')
        data_font = Font(name='微软雅黑', size=10)
        leader_font = Font(name='微软雅黑', size=10, bold=True)

        # 大标题样式
        big_title_font = Font(name='微软雅黑', size=16, bold=True, color='FFFFFF')
        big_title_fill = PatternFill(start_color='2d3748', end_color='2d3748', fill_type='solid')

        # 分数标色字体（<60绿，>150红）
        red_font = Font(name='微软雅黑', size=10, color='dc2626', bold=True)
        green_font = Font(name='微软雅黑', size=10, color='16a34a', bold=True)

        title_fill = PatternFill(start_color='4a5568', end_color='4a5568', fill_type='solid')
        header_fill = PatternFill(start_color='e2e8f0', end_color='e2e8f0', fill_type='solid')
        leader_fill = PatternFill(start_color='f7fafc', end_color='f7fafc', fill_type='solid')

        thin_border = Border(
            left=Side(style='thin', color='cbd5e0'),
            right=Side(style='thin', color='cbd5e0'),
            top=Side(style='thin', color='cbd5e0'),
            bottom=Side(style='thin', color='cbd5e0')
        )

        center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

        # ========== 第1行：大标题（合并10列） ==========
        ws.merge_cells('A1:J1')
        title_cell = ws.cell(row=1, column=1)
        title_cell.value = big_title
        title_cell.font = big_title_font
        title_cell.fill = big_title_fill
        title_cell.alignment = center_align

        for col in range(1, 11):
            ws.cell(row=1, column=col).border = thin_border

        # ========== 第2行：一级表头 ==========
        ws.merge_cells('A2:A3')  # 序号
        ws.merge_cells('B2:B3')  # 姓名
        ws.merge_cells('C2:C3')  # 岗位
        ws.merge_cells('D2:D3')  # 考核成绩
        ws.merge_cells('E2:G2')  # 业绩订单销售量（3列）
        ws.merge_cells('H2:J2')  # 考核得分项（3列）

        headers_level1 = ['序号', '姓名', '岗位', '考核成绩', '业绩订单销售量', '考核得分项']
        for col, header in enumerate(headers_level1, 1):
            if col == 5:
                cell = ws.cell(row=2, column=5)
            elif col == 6:
                cell = ws.cell(row=2, column=8)
            else:
                cell = ws.cell(row=2, column=col)
            cell.value = header
            cell.font = title_font
            cell.fill = title_fill
            cell.alignment = center_align
            cell.border = thin_border

        # ========== 第3行：二级表头 ==========
        headers_level2 = ['KPI订单量', '完成订单量', '是否达标',
                          '业绩指标得分', '行为考核得分', '加分项得分']
        for col, header in enumerate(headers_level2, 5):
            cell = ws.cell(row=3, column=col)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

        for row in [2, 3]:
            for col in range(1, 11):
                ws.cell(row=row, column=col).border = thin_border

        # ========== 填充数据（从第4行开始） ==========
        row = 4
        idx = 1
        for item in data_list:
            is_leader_row = item['is_leader_row']
            final_score_val = item['final_score']

            data = [
                idx,
                item['employee_name'],
                item['role'],
                final_score_val,
                item['target_orders'],
                item['actual_orders'],
                item['is_achieved'],
                item['performance_score'],
                item['behavior_score'],
                item['bonus_score'],
            ]

            for col, value in enumerate(data, 1):
                cell = ws.cell(row=row, column=col)
                cell.value = value
                cell.alignment = center_align
                cell.border = thin_border

                if is_leader_row:
                    cell.font = leader_font
                    cell.fill = leader_fill
                else:
                    cell.font = data_font

                if col == 4 and isinstance(final_score_val, (int, float, Decimal)):
                    if final_score_val < 60:
                        cell.font = green_font
                    elif final_score_val > 150:
                        cell.font = red_font

                if col == 7 and value == '是':
                    cell.font = red_font

            row += 1
            idx += 1

        # 设置列宽
        column_widths = [6, 12, 12, 12, 12, 12, 10, 12, 12, 12]
        for i, width in enumerate(column_widths, 1):
            ws.column_dimensions[chr(64 + i)].width = width

        # 设置行高
        ws.row_dimensions[1].height = 40
        ws.row_dimensions[2].height = 30
        ws.row_dimensions[3].height = 35

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{file_name}"'
        return response
