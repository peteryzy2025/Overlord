from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from django.views.decorators.http import require_http_methods
from django.http import JsonResponse, HttpResponse
from django.db import transaction
from django.db.models import Q
from django.core.paginator import Paginator
import json
import openpyxl

from general.models import User
from task.models import ProductRequirement
from task.utils import parse_permissions
from api.wc.crawler_wc import get_ykartwood_product

@login_required
def product_create_page(request):
    """
    产品需求创建页面
    """
    return render(request, 'product/product_create.html')

@login_required
def product_list_page(request):
    """
    产品需求列表页面
    """
    return render(request, 'product/product_list.html')

@login_required
@require_http_methods(["POST"])
@transaction.atomic
def create_product_requirement_api(request):
    """
    创建产品需求 (简化版)
    """
    try:
        data = json.loads(request.body)
        
        # 必填字段验证
        required_fields = ['platform', 'product_url']
        
        for field in required_fields:
            if not data.get(field):
                return JsonResponse({'success': False, 'message': f'字段 {field} 不能为空'}, status=400)
                
        # 创建对象
        requirement = ProductRequirement.objects.create(
            created_by=request.user,
            platform=data.get('platform'),
            product_url=data.get('product_url'),
            status=ProductRequirement.STATUS_SUBMITTED
        )
        
        return JsonResponse({
            'success': True,
            'data': {'id': requirement.id, 'redirect_url': '/task/product/list/'}
        })
        
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)}, status=500)


@login_required
@require_http_methods(["GET"])
def export_product_requirement_excel_api(request, pk):
    """
    导出产品需求为Excel
    GET /api/products/<pk>/export/
    """
    try:
        item = ProductRequirement.objects.get(pk=pk)
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "产品详情"
        
        # Helper to write row
        row_num = 1
        def write_row(label, value):
            nonlocal row_num
            ws.cell(row=row_num, column=1, value=label)
            ws.cell(row=row_num, column=2, value=str(value) if value is not None else '')
            row_num += 1
            
        # Basic Info
        ws.merge_cells('A1:B1')
        ws['A1'] = "基本信息"
        ws['A1'].font = openpyxl.styles.Font(bold=True)
        row_num += 1
        
        write_row('需求标题', item.task.title if item.task else '-') # item.title removed
        write_row('需求单号', item.requirement_no or (item.task.task_no if item.task else '-'))
        write_row('外采平台', item.get_platform_display())
        write_row('上架平台', item.get_listing_platform_display())
        write_row('产品链接', item.product_url)
        write_row('产品ID', item.product_id)
        write_row('产品名称', item.product_name)
        write_row('产品简称', item.product_abbr)
        write_row('英文名称', item.english_name)
        write_row('状态', item.get_status_display())
        write_row('材质', item.material)
        write_row('生产工艺', item.craft)
        write_row('计量单位', item.unit)
        
        # Design Info
        row_num += 1
        ws.merge_cells(f'A{row_num}:B{row_num}')
        ws[f'A{row_num}'] = "设计说明"
        ws[f'A{row_num}'].font = openpyxl.styles.Font(bold=True)
        row_num += 1
        
        write_row('设计说明', item.design_desc)
        write_row('设计区域', item.design_area)
        write_row('图片要求', item.image_requirement)
        write_row('备注', item.remark)
        
        # Customs Info
        row_num += 1
        ws.merge_cells(f'A{row_num}:B{row_num}')
        ws[f'A{row_num}'] = "报关信息"
        ws[f'A{row_num}'].font = openpyxl.styles.Font(bold=True)
        row_num += 1
        
        write_row('中文名称', item.customs_cn_name)
        write_row('英文名称', item.customs_en_name)
        write_row('申报重量(g)', item.declared_weight)
        write_row('申报单价($)', item.declared_price)
        write_row('海关编码', item.customs_code)
        write_row('材质(中)', item.material_cn)
        write_row('材质(英)', item.material_en)
        write_row('商标类目', item.trademark_category)
        write_row('所属分类', item.category)
        write_row('特殊货物类型', item.special_cargo_type)
        write_row('产品标识', item.product_label)
        
        # Other Desc
        row_num += 1
        ws.merge_cells(f'A{row_num}:B{row_num}')
        ws[f'A{row_num}'] = "其他描述"
        ws[f'A{row_num}'].font = openpyxl.styles.Font(bold=True)
        row_num += 1
        
        write_row('材质说明', item.material_desc)
        write_row('配件构造', item.accessory_struct)
        write_row('产品性能', item.product_performance)
        write_row('适用情景', item.applicable_scenario)
        write_row('洗涤说明', item.washing_instructions)
        write_row('特别说明', item.special_note)
        write_row('温馨提醒', item.reminder)
        
        # Extra Info
        row_num += 1
        ws.merge_cells(f'A{row_num}:B{row_num}')
        ws[f'A{row_num}'] = "额外信息"
        ws[f'A{row_num}'].font = openpyxl.styles.Font(bold=True)
        row_num += 1
        
        # Colors
        colors = item.color_name
        if isinstance(colors, list):
            colors = ", ".join(colors)
        write_row('颜色列表', colors)

        # Sizes
        sizes = item.size_list
        if isinstance(sizes, list):
            sizes = ", ".join(sizes)
        write_row('尺码列表', sizes)
        
        # Packaging
        row_num += 1
        ws.merge_cells(f'A{row_num}:B{row_num}')
        ws[f'A{row_num}'] = "包装规格"
        ws[f'A{row_num}'].font = openpyxl.styles.Font(bold=True)
        row_num += 1
        
        specs = item.packaging_specification
        if specs and isinstance(specs, list):
            # Write table
            for r_idx, row_data in enumerate(specs):
                for c_idx, cell_data in enumerate(row_data):
                    content = cell_data.get('content', '')
                    remark = cell_data.get('remark', '')
                    val = content
                    if remark:
                        val += f"\n({remark})"
                    ws.cell(row=row_num, column=c_idx+1, value=val)
                row_num += 1
        else:
            write_row('包装规格', '暂无数据')

        # Adjust column width
        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 50

        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        filename = request.GET.get('filename', item.product_name or 'product')
        import urllib.parse
        filename = urllib.parse.quote(filename)
        response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"; filename*=UTF-8\'\'{filename}.xlsx'
        
        wb.save(response)
        return response
        
    except ProductRequirement.DoesNotExist:
        return JsonResponse({'success': False, 'message': '未找到该记录'}, status=404)
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)}, status=500)

@login_required
@require_http_methods(["GET"])
def get_product_requirements_api(request):
    """
    获取产品需求列表 (API)
    GET /api/products/list/
    """
    try:
        page = int(request.GET.get('page', 1))
        page_size = int(request.GET.get('page_size', 20))
        search = request.GET.get('search', '').strip()
        
        queryset = ProductRequirement.objects.all().order_by('-created_at')
        
        # 筛选
        status = request.GET.get('status')
        if status:
            status_list = status.split(',')
            queryset = queryset.filter(status__in=status_list)

        created_by = request.GET.get('created_by')
        if created_by:
            created_by_list = created_by.split(',')
            queryset = queryset.filter(created_by_id__in=created_by_list)
            
        if search:
            queryset = queryset.filter(
                Q(product_name__icontains=search) | 
                # Q(title__icontains=search) | 
                Q(requirement_no__icontains=search) |
                Q(product_id__icontains=search) # 支持搜索 Product ID
            )
            
        paginator = Paginator(queryset, page_size)
        page_obj = paginator.get_page(page)
        
        data = []
        for item in page_obj:
            data.append({
                'id': item.id,
                'title': item.task.title if item.task else '-', # item.title removed
                'requirement_no': item.requirement_no or (item.task.task_no if item.task else '-'),
                'product_url': item.product_url,
                'product_id': item.product_id, # 返回 Product ID
                'platform': item.get_platform_display(), # 返回平台名称
                'platform_code': item.platform, # 返回平台代码，供前端逻辑使用
                'status': item.status,
                'status_display': item.get_status_display(),
                'created_by': item.created_by.first_name or item.created_by.username,
                'created_at': item.created_at.strftime('%Y-%m-%d %H:%M:%S'),
            })
            
        return JsonResponse({
            'success': True, 
            'data': {
                'items': data,
                'page': page,
                'total': paginator.count,
                'total_pages': paginator.num_pages
            }
        })
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)}, status=500)

@login_required
@require_http_methods(["POST"])
def recrawl_product_requirement_api(request):
    """
    重新爬取产品需求 (仅限艺之冠)
    POST /api/products/recrawl/
    Body: { "id": 1, "product_id": "123", "platform": "yizhiguan" }
    """
    try:
        data = json.loads(request.body)
        req_id = data.get('id')
        product_id = data.get('product_id')
        platform = data.get('platform')
        
        if not req_id or not product_id:
            return JsonResponse({'success': False, 'message': '参数缺失'}, status=400)
            
        if platform != 'yizhiguan':
            return JsonResponse({'success': False, 'message': '仅支持艺之冠平台重新爬取'}, status=400)
            
        # 获取对象
        try:
            item = ProductRequirement.objects.get(pk=req_id)
        except ProductRequirement.DoesNotExist:
             return JsonResponse({'success': False, 'message': '记录不存在'}, status=404)
             
        # 调用爬虫
        try:
            crawler_data = get_ykartwood_product(product_id, item.listing_platform, item.platform)
            if not crawler_data:
                return JsonResponse({'success': False, 'message': '爬取失败，未获取到数据'}, status=500)
                
            # 更新字段
            item.product_name = crawler_data.get('product_name', item.product_name)
            item.product_abbr = crawler_data.get('product_abbr', item.product_abbr)
            item.english_name = crawler_data.get('english_name', item.english_name)
            item.material = crawler_data.get('material', item.material)
            item.craft = crawler_data.get('craft', item.craft)
            item.unit = crawler_data.get('unit', item.unit)
            
            item.customs_cn_name = crawler_data.get('customs_cn_name', item.customs_cn_name)
            item.customs_en_name = crawler_data.get('customs_en_name', item.customs_en_name)
            item.declared_weight = crawler_data.get('declared_weight') # 整数，直接覆盖
            item.declared_price = crawler_data.get('declared_price') # 浮点数，直接覆盖
            
            item.material_cn = crawler_data.get('material_cn', item.material_cn)
            item.material_desc = crawler_data.get('material_desc', item.material_desc)
            item.accessory_struct = crawler_data.get('accessory_struct', item.accessory_struct)
            item.product_performance = crawler_data.get('product_performance', item.product_performance)
            item.applicable_scenario = crawler_data.get('applicable_scenario', item.applicable_scenario)
            item.washing_instructions = crawler_data.get('washing_instructions', item.washing_instructions)
            item.special_note = crawler_data.get('special_note', item.special_note)
            item.reminder = crawler_data.get('reminder', item.reminder)
            item.design_desc = crawler_data.get('design_desc', item.design_desc)
            item.design_area = crawler_data.get('design_area', item.design_area)
            
            # 新增字段
            item.color_name = crawler_data.get('color_name', item.color_name)
            item.img_urls_list = crawler_data.get('img_urls_list', item.img_urls_list)
            
            # 新增报关字段
            item.category = crawler_data.get('category', item.category)
            item.special_cargo_type = crawler_data.get('special_cargo_type', item.special_cargo_type)
            item.product_label = crawler_data.get('product_label', item.product_label)
            item.packaging_specification = crawler_data.get('packaging_specification', item.packaging_specification)
            item.product_size = crawler_data.get('product_size', item.product_size)
            
            item.save()
            
            return JsonResponse({'success': True, 'message': '重新爬取并更新成功'})
            
        except Exception as e:
            return JsonResponse({'success': False, 'message': f'爬虫执行出错: {str(e)}'}, status=500)
            
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)}, status=500)


@login_required
@require_http_methods(["POST"])
def reject_product_requirements_api(request):
    """
    批量驳回产品需求
    POST /api/products/reject/
    Body: { "ids": [1, 2, 3] }
    """
    try:
        # 权限验证：只有拥有 555 权限码的用户可以驳回
        permissions = parse_permissions(getattr(request.user, 'permission', ''))
        
        if '555' not in permissions:
             return JsonResponse({'success': False, 'message': '无权进行驳回操作'}, status=403)

        data = json.loads(request.body)
        ids = data.get('ids', [])
        
        if not ids:
            return JsonResponse({'success': False, 'message': '请选择要驳回的记录'}, status=400)
            
        # 批量更新状态
        ProductRequirement.objects.filter(id__in=ids).update(status=ProductRequirement.STATUS_REJECTED)
        
        return JsonResponse({'success': True, 'message': f'成功驳回 {len(ids)} 条记录'})
        
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)}, status=500)

@login_required
@require_http_methods(["GET"])
def get_product_requirement_detail_api(request, pk):
    """
    获取产品需求详情 (API)
    GET /api/products/<pk>/
    """
    try:
        item = ProductRequirement.objects.get(pk=pk)
        
        # 权限检查：是否可以编辑
        # 只有创建者、所有者、组长或管理员可以编辑
        can_edit = False
        current_user = request.user
        if current_user == item.created_by or current_user == item.owner:
            can_edit = True
        else:
            permissions = parse_permissions(getattr(current_user, 'permission', ''))
            if 'ops_all' in permissions:
                can_edit = True
            elif 'ops_group' in permissions and hasattr(current_user, 'operational_account'):
                # 检查组权限
                if item.created_by.operational_account and \
                   item.created_by.operational_account.ops_group == current_user.operational_account.ops_group:
                    can_edit = True
                    
        data = {
            'id': item.id,
            'title': item.task.title if item.task else '-', # item.title removed
            'requirement_no': item.requirement_no or (item.task.task_no if item.task else '-'),
            'platform': item.platform,
            'platform_display': item.get_platform_display(), # 增加显示名称
            'listing_platform': item.listing_platform, # 上架平台
            'product_url': item.product_url,
            'product_id': item.product_id, # Product ID
            'product_name': item.product_name,
            'product_abbr': item.product_abbr,
            'english_name': item.english_name,
            'status': item.status,
            'status_display': item.get_status_display(),
            'material': item.material,
            'craft': item.craft,
            'unit': item.unit,
            
            # 报关信息
            'customs_cn_name': item.customs_cn_name,
            'customs_en_name': item.customs_en_name,
            'declared_weight': str(item.declared_weight) if item.declared_weight else '',
            'declared_price': str(item.declared_price) if item.declared_price else '',
            'customs_code': item.customs_code,
            'material_cn': item.material_cn,
            'material_en': item.material_en,
            
            # 描述字段
            'material_desc': item.material_desc,
            'accessory_struct': item.accessory_struct,
            'product_performance': item.product_performance,
            'applicable_scenario': item.applicable_scenario,
            'washing_instructions': item.washing_instructions,
            'special_note': item.special_note,
            'reminder': item.reminder,
            'design_desc': item.design_desc,
            'design_area': item.design_area,
            'image_requirement': item.image_requirement,
            'remark': item.remark,
            'status': item.status,
            'status_display': item.get_status_display(),
            
            # 包装信息
            'packaging_specification': item.packaging_specification,

            # 报关额外信息
            'trademark_category': item.trademark_category,
            'category': item.category,
            'special_cargo_type': item.special_cargo_type,
            'product_label': item.product_label,

            # 额外信息
            'color_name': item.color_name,
            'size_list': item.size_list,
            'img_urls_list': item.img_urls_list,
            'product_size': item.product_size,
        }
        
        return JsonResponse({
            'success': True,
            'data': data,
            'meta': {
                'can_edit': can_edit
            }
        })
        
    except ProductRequirement.DoesNotExist:
        return JsonResponse({'success': False, 'message': '需求不存在'}, status=404)
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)}, status=500)

@login_required
@require_http_methods(["POST"])
def update_product_requirement_api(request, pk):
    """
    更新产品需求 (仅限拥有 555 权限的用户)
    """
    try:
        # 权限检查
        permissions = parse_permissions(getattr(request.user, 'permission', ''))
        if '555' not in permissions:
            return JsonResponse({'success': False, 'message': '无权限修改'}, status=403)
            
        item = ProductRequirement.objects.get(pk=pk)
        data = json.loads(request.body)
        
        # 更新字段
        fields = [
            'platform', 'listing_platform', 'product_url',
            'product_name', 'product_abbr', 'english_name', 'material',  
            'craft', 'unit', 'customs_cn_name', 'customs_en_name',
            'declared_weight', 'declared_price', 'customs_code',
            'material_cn', 'material_en', 'material_desc', 
            'accessory_struct', 'product_performance', 'applicable_scenario',
            'washing_instructions', 'special_note', 'reminder',
            'design_desc', 'design_area', 'image_requirement', 'remark', 'status',
            'color_name', 'size_list', 'img_urls_list',
            'category', 'special_cargo_type', 'product_label', 'trademark_category',
            'packaging_specification', 'product_size'
        ]
        
        for field in fields:
            if field in data:
                # 特殊处理 JSON 字段
                if field in ['color_name', 'img_urls_list', 'packaging_specification', 'size_list', 'product_size']:
                     try:
                        # 如果是字符串，尝试解析为 JSON
                         if isinstance(data[field], str):
                             import json
                             setattr(item, field, json.loads(data[field]))
                         else:
                             setattr(item, field, data[field])
                     except:
                         # 解析失败则跳过或设为空列表，视需求而定
                         pass
                else:
                    setattr(item, field, data[field])
        
        item.save()
        
        return JsonResponse({'success': True, 'message': '更新成功'})
        
    except ProductRequirement.DoesNotExist:
        return JsonResponse({'success': False, 'message': '未找到该记录'}, status=404)
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)}, status=500)
