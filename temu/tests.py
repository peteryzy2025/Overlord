def print_tracking_column(file_path):
    """打印Excel中'物流单号'列的所有数据"""
    import openpyxl

    # 读取文件
    workbook = openpyxl.load_workbook(file_path, data_only=True)
    worksheet = workbook.active

    # 查找列位置
    target_col = None
    for idx, cell in enumerate(worksheet[1], 1):
        if cell.value and "物流单号" in str(cell.value):
            target_col = idx
            break

    if not target_col:
        print("未找到'物流单号'列")
        return

    # 打印该列所有值
    print(f"\n--- {worksheet.cell(row=1, column=target_col).value} ---")
    for row in range(2, worksheet.max_row + 1):
        value = worksheet.cell(row=row, column=target_col).value
        if value:
            print(value)

    workbook.close()


# 直接使用
print_tracking_column(r"D:\下载目录\导出交易记录_1765847658858.xlsx")